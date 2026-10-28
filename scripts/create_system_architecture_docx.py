"""Create a VISHUE system-architecture submission document from the provided template."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = Path(r"C:\Users\playdata2\Downloads\[모델배포] 시스템 구성도_양식.docx")
OUTPUT = ROOT / "docs" / "deliverables" / "VISHUE_시스템구성도_제출본.docx"
DIAGRAM = ROOT / "docs" / "deliverables" / "vishue-system-architecture.png"


def clear_document(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_font(run, size: int, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(document: Document, value: str, size: int = 10, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    run = paragraph.add_run(value)
    set_font(run, size, bold)


def add_heading(document: Document, value: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(value)
    set_font(run, 16 if level == 1 else 12, True, "7A1F3D")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "D6C5CC")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_font(run, 9, True, "FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "7A1F3D")
        cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(value)
            set_font(run, 8)
    document.add_paragraph()


def main() -> None:
    template = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_TEMPLATE
    if not template.exists():
        raise FileNotFoundError(template)
    if not DIAGRAM.exists():
        raise FileNotFoundError(DIAGRAM)

    document = Document(template)
    clear_document(document)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    add_text(document, "SK 네트웍스 Family AI 00기 : [작성 필요]  |  3__팀", 10)
    add_text(document, "VISHUE 시스템 구성도", 23, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "산출물 단계: 모델배포  |  제출 일자: [작성 필요]  |  GitHub 경로: [작성 필요]  |  작성 팀원: [작성 필요]", 9, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(document, "1. 전체 서비스 구조 (Service Architecture)")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(10.4))
    add_text(document, "※ 운영 EKS·RDS·ElastiCache는 배포 예정 설계이며, 개발 환경은 Render + Supabase + Neo4j AuraDB Free를 사용합니다.", 8)

    add_heading(document, "2. 시스템 구성 요소 상세 (System Components)")
    add_table(
        document,
        ["구분", "구성 요소", "책임 / 연결", "상태"],
        [
            ["클라이언트", "Expo React Native 앱", "룩·제품·구매·챗봇 화면, API 호출", "구현됨"],
            ["기기 내 AR", "Camera + WebView + MediaPipe", "얼굴 랜드마크·가상 메이크업·피부톤 수치 산출", "구현됨"],
            ["인증", "Supabase Auth", "이메일 코드·Google 로그인, JWT 발급", "구현됨"],
            ["API", "FastAPI (vishue-api)", "카탈로그·셰이드 추천·YouTube/네이버 연동, PostgreSQL 소유", "구현됨"],
            ["챗봇", "FastAPI (vishue-chatbot)", "Neo4j GraphRAG, 외부 LLM API 호출", "구현됨"],
            ["데이터", "PostgreSQL / Redis / Neo4j", "원본 카탈로그 / 캐시 / 챗봇 그래프 사본", "구현됨·운영 예정"],
            ["운영", "AWS EKS + RDS + ElastiCache", "Terraform · Helm · Argo CD GitOps 배포", "운영 예정"],
        ],
        [1.1, 2.4, 5.9, 1.0],
    )

    add_heading(document, "3. 연결 관계 및 데이터 흐름")
    add_text(document, "① 얼굴 영상·사진은 기기 밖으로 전송하지 않고 WebView/MediaPipe에서 처리합니다. ② 앱은 Supabase access token을 붙여 Gateway의 HTTPS API를 호출합니다. ③ api만 PostgreSQL 원본을 읽고 쓰며, sync-graph가 API HTTP 계약을 통해 Neo4j를 단방향 갱신합니다. ④ chatbot은 Neo4j와 외부 LLM API를 사용하며 PostgreSQL에 직접 접속하지 않습니다.", 9)

    add_heading(document, "4. 기술 스택 및 버전 명세 (Tech Stack & Versions)")
    add_table(
        document,
        ["구분", "기술", "버전", "용도"],
        [
            ["모바일", "Expo / React Native", "^58.0.19 / 0.88.0", "모바일·웹 앱 런타임"],
            ["UI·상태", "React / Expo Router / React Query / Zustand", "19.3.0 / ~58.0.26 / ^5.104.3 / ^5.0.15", "화면·라우팅·서버 상태·UI 상태"],
            ["AR", "Camera + WebView + MediaPipe Face Landmarker", "Expo SDK 58 기준", "기기 내 얼굴 랜드마크·메이크업 렌더링"],
            ["백엔드", "Python / FastAPI / Uvicorn", ">=3.12 / >=0.161.1 / >=0.54.0", "비동기 API·챗봇 서비스"],
            ["DB", "PostgreSQL / SQLAlchemy async / asyncpg / Alembic", "17 / >=2.1.6 / >=0.32.0 / >=1.20.1", "카탈로그 원본·마이그레이션"],
            ["그래프·캐시", "Neo4j / redis-py", "Neo4j 5.x / >=8.1.0", "GraphRAG·캐시·레이트 리밋"],
            ["배포", "Docker / Kubernetes / Helm / Argo CD / Terraform", "인프라 관리 버전 기준", "컨테이너·GitOps·AWS 인프라"],
        ],
        [1.1, 3.4, 2.6, 3.3],
    )

    add_heading(document, "5. 보안 및 운영 관리")
    add_text(document, "[확인] API 키와 DB 연결값은 .env(로컬) 또는 AWS Secrets Manager(운영)에서 관리합니다.  [확인] 얼굴 영상·사진은 서버로 전송·저장하지 않습니다.  [확인] Supabase JWT를 백엔드에서 JWKS로 검증합니다.  [확인] /healthz와 /readyz를 분리하고 Redis 장애 시 제한적으로 동작합니다.", 9)
    add_text(document, "미도입 항목: LangGraph, MCP Server, 자체 GPU/파인튜닝 모델 서빙. 현재 구조에 존재하지 않는 구성요소는 다이어그램에 포함하지 않았습니다.", 8)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
